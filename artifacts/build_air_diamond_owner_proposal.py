from decimal import Decimal as D, ROUND_HALF_UP
from pathlib import Path
import json

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs/air_diamond_owner_proposal'
OUT.mkdir(parents=True, exist_ok=True)
FONT = 'TH Sarabun New'
CAPEX = D('4050000')
KWH = D('27')
PRICE = D('7.90')
POWER = D('4.00')
CPO = D('.08')
OM = D('36000')


def model(cars):
    energy = D(cars) * KWH * 365
    revenue = energy * PRICE
    platform = revenue * CPO
    electricity = energy * POWER
    cash = revenue - platform - electricity - OM
    return dict(energy=energy, revenue=revenue, platform=platform,
                electricity=electricity, om=OM, cash=cash, payback=CAPEX / cash)


def money(n):
    return f'{D(n).quantize(D(1), rounding=ROUND_HALF_UP):,}'


def font(obj, size, bold=False):
    obj.font.name = FONT
    obj.font.size = Pt(size)
    obj.font.bold = bold
    obj.font.color.rgb = RGBColor(0, 0, 0)
    rp = obj._element.get_or_add_rPr()
    rf = rp.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rp.insert(0, rf)
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rf.set(qn('w:' + key), FONT)
    for key in ('sz', 'szCs'):
        el = rp.find(qn('w:' + key))
        if el is None:
            el = OxmlElement('w:' + key)
            rp.append(el)
        el.set(qn('w:val'), str(int(size * 2)))


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.65), Cm(2.15)
sec.left_margin = sec.right_margin = Cm(1.9)
sec.footer_distance = Cm(.7)
for name, size, bold in [('Normal', 16, False), ('Title', 26, True),
                         ('Subtitle', 19, False), ('Heading 1', 20, True),
                         ('Heading 2', 17, True)]:
    style = doc.styles[name]
    font(style, size, bold)
    style.paragraph_format.space_after = Pt(3)
    style.paragraph_format.line_spacing = 1.03
    if name.startswith('Heading'):
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.keep_with_next = True
for style in doc.styles:
    pp = style.element.pPr
    if pp is not None:
        for el in list(pp.findall(qn('w:pBdr'))):
            pp.remove(el)


def para(text, bold=False, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    font(r, size or 16, bold)
    return p


def heading(text, newpage=False):
    p = doc.add_paragraph(text, 'Heading 1')
    p.paragraph_format.page_break_before = newpage
    return p


def table(headers, rows, widths, size=14.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, w in zip(t.columns, widths):
        col.width = Cm(w)
    for cell, txt in zip(t.rows[0].cells, headers):
        cell.text = txt
    for row in rows:
        for cell, txt in zip(t.add_row().cells, row):
            cell.text = str(txt)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            el.set(qn('w:' + k), v)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for ri, row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if ri == 0:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:fill'), 'EDEDED' if ri == 0 else ('FAFAFA' if ri % 2 else 'FFFFFF'))
            cp.append(shd)
            mar = OxmlElement('w:tcMar')
            for edge, value in [('top', '45'), ('bottom', '45'), ('left', '95'), ('right', '95')]:
                el = OxmlElement('w:' + edge)
                el.set(qn('w:w'), value)
                el.set(qn('w:type'), 'dxa')
                mar.append(el)
            cp.append(mar)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                if ri == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    font(r, size, ri == 0)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(3)
    font(p.add_run(), 2)
    return t


def link(label, url):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    relation = doc.part.relate_to(url,
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), relation)
    r = OxmlElement('w:r')
    rp = OxmlElement('w:rPr')
    rf = OxmlElement('w:rFonts')
    for k in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rf.set(qn('w:' + k), FONT)
    rp.append(rf)
    for k, v in [('sz', '26'), ('szCs', '26'), ('color', '000000'), ('u', 'single')]:
        el = OxmlElement('w:' + k)
        el.set(qn('w:val'), v)
        rp.append(el)
    r.append(rp)
    txt = OxmlElement('w:t')
    txt.text = label
    r.append(txt)
    h.append(r)
    p._p.append(h)


foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(foot.add_run('Air Diamond   |   TCE Project   |   '), 11)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
foot._p.append(field)

m = model(48)
logo = Path('C:/Users/Lenovo/Desktop/logo.png')
if logo.exists():
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.add_run().add_picture(str(logo), width=Cm(3.0))
doc.add_paragraph('ข้อเสนอการลงทุนสถานีชาร์จรถยนต์ไฟฟ้า', 'Title')
doc.add_paragraph('Air Diamond Cafe and Hotel แม่แตง', 'Subtitle')
para('เสนอเจ้าของพื้นที่ 111 หมู่ 3 ตำบลขี้เหล็ก อำเภอแม่แตง จังหวัดเชียงใหม่\nแยกบายพาสแม่มาลัยไปปาย   พิกัด 19.072965, 98.938715', size=14)
para('TCE Project   |   23 ตุลาคม 2569', size=13)
para('TCE Project เสนอให้พัฒนาสถานีชาร์จ DC กำลังรวม 360 kW ด้วยตู้ 180 kW จำนวน 2 ตู้ 4 ช่องจอด เพื่อเพิ่มบริการให้ลูกค้าคาเฟ่ โรงแรม และผู้เดินทางที่เลือกแวะพัก พร้อมสร้างรายได้ค่าชาร์จให้เจ้าของพื้นที่โดยตรง')
para('เงินลงทุน 4.05 ล้านบาท ตั้งเป้าการใช้บริการเฉลี่ย 48 คันต่อวัน หรือ 12 คันต่อช่องจอดต่อวัน ภายใต้ยอดใช้งานและต้นทุนตามแผน จะคืนทุนประมาณ 2.7 ปี', True, 18)
heading('ภาพรวมการลงทุน')
table(['รายการ', 'ประมาณการตามเป้าหมายการใช้บริการ'], [
    ('รูปแบบสถานี', 'DC 180 kW จำนวน 2 ตู้ 4 ช่องจอด'),
    ('เงินลงทุนเริ่มต้น', '4,050,000 บาท'),
    ('เป้าหมายรถเข้าชาร์จเฉลี่ย', '48 คันต่อวัน'),
    ('รายได้ค่าชาร์จเฉลี่ยต่อเดือน', money(m['revenue'] / 12) + ' บาท'),
    ('เงินสดหลังค่าใช้จ่ายเฉลี่ยต่อเดือน', money(m['cash'] / 12) + ' บาท'),
    ('คืนทุนจากกระแสเงินสดดำเนินงาน', 'ประมาณ 2 ปี 8 เดือน'),
], [8.1, 9.1], 15)
para('ฐานรายได้ใช้พลังงานเฉลี่ย 27 kWh ต่อคัน ราคาขาย 7.90 บาทต่อหน่วย และไม่มีค่าเช่าหรือ GP เจ้าของพื้นที่ เนื่องจากเจ้าของพื้นที่ลงทุนเอง', size=14)
heading('โอกาสที่สถานีสร้างให้กับพื้นที่')
para('จุดหมายท่องเที่ยวธีมเครื่องบิน คาเฟ่ และโรงแรม ช่วยให้ลูกค้ามีกิจกรรมระหว่างรอชาร์จ ส่วนตลาดแม่มาลัยและย่านการค้าใกล้เคียงเสริมบริบทของทำเล การเชื่อมบริการชาร์จกับอาหาร เครื่องดื่ม และที่พัก จึงเป็นแนวทางเพิ่มเหตุผลให้ผู้ใช้รถ EV เลือกแวะโครงการ')
para('48 คันต่อวันเป็นเป้าหมายเชิงพาณิชย์ที่ใช้ประเมินการลงทุน ไม่ใช่ยอดลูกค้าที่ตรวจนับหรือดีมานด์ที่ยืนยันแล้ว ต้องตรวจจำนวนรถเข้าโครงการและสัดส่วนผู้เลือกชาร์จก่อนตัดสินใจลงทุน', size=14)

heading('รายได้และเงินสดที่เจ้าของพื้นที่ได้รับ', True)
para('คำนวณเมื่อสถานีให้บริการได้เฉลี่ย 48 คันต่อวันตลอดปี รายได้ในตารางนับเฉพาะการขายไฟชาร์จรถ ไม่รวมรายได้เพิ่มจากคาเฟ่ โรงแรม หรือร้านค้า')
table(['รายการ', 'เฉลี่ยต่อเดือน บาท', 'ต่อปี บาท'], [
    ('รายได้ค่าชาร์จ', money(m['revenue'] / 12), money(m['revenue'])),
    ('หักต้นทุนไฟฟ้า', money(m['electricity'] / 12), money(m['electricity'])),
    ('หักค่าระบบ CPO 8% ของรายได้', money(m['platform'] / 12), money(m['platform'])),
    ('หักค่าดูแลสถานี O&M', money(OM / 12), money(OM)),
    ('เงินสดจากการดำเนินงาน', money(m['cash'] / 12), money(m['cash'])),
], [8, 4.6, 4.6], 14.5)
para('ตัวเลขต่อเดือนเป็นค่าเฉลี่ยทั้งปี ไม่ใช่ประมาณการแยกฤดูกาล เงินสดจากการดำเนินงานคือรายได้หลังหักรายการข้างต้น ยังไม่ใช่กำไรสุทธิทางบัญชี', size=13.5)
heading('ประมาณการเงินสดสะสม 10 ปี')
para('แสดงปีดำเนินงานนับจากเปิดสถานี โดยคงยอด 48 คันต่อวันและราคาต้นทุนเดิมทุกปี ไม่เพิ่มยอดรถเพื่อเร่งการคืนทุน', size=14)
rows = [('ก่อนเปิด', '-', '-', '-4,050,000')]
for year in range(1, 11):
    rows.append((str(year), money(m['revenue']), money(m['cash']), money(m['cash'] * year - CAPEX)))
table(['ปีดำเนินงาน', 'รายได้ต่อปี บาท', 'เงินสดดำเนินงานต่อปี บาท', 'เงินสดสะสมหักลงทุน บาท'],
      rows, [2.2, 4.7, 5.0, 5.3], 14)
para('ครบ 10 ปี รายได้รวม 37.37 ล้านบาท เงินสดจากการดำเนินงานรวม 15.10 ล้านบาท และเงินสดสะสมหลังหักเงินลงทุนเริ่มต้น 11.05 ล้านบาท', True, 15)
para('คืนทุนแบบไม่คิดลดประมาณ 2.68 ปี คำนวณจากเงินลงทุนหารเงินสดดำเนินงานต่อปี โดยสมมติว่าเฉลี่ยยอดใช้งานได้ตามแผนตั้งแต่ปีแรก หากมีช่วงเริ่มต้นสะสมลูกค้า ระยะเวลาคืนทุนจะยาวขึ้น', size=13.5)

heading('รูปแบบบริการและเงื่อนไขสู่เป้าหมาย', True)
para('สถานี 4 ช่องจอดรองรับรถที่เข้าพร้อมกัน และยังเหลือ 2 ช่องเมื่อตู้หนึ่งซ่อมบำรุง ทั้งนี้ 180 kW เป็นกำลังรวมของแต่ละตู้ การใช้สองหัวพร้อมกันจะแบ่งกำลังตามระบบและความสามารถของรถ ไม่ใช่ 180 kW ต่อช่องพร้อมกัน')
para('พื้นที่เข้าออกได้ตลอด 24 ชั่วโมงจากการมีธุรกิจโรงแรม ควรจัดช่องชาร์จเฉพาะ ป้ายจากถนน แสงสว่าง และช่องทางช่วยเหลือให้พร้อมตามเวลาเปิดบริการ')
heading('สร้างลูกค้าร่วมกับคาเฟ่และโรงแรม')
para('เชื่อมการชาร์จกับคูปองอาหาร เครื่องดื่ม หรือสิทธิพิเศษที่พัก แสดงสถานะหัวชาร์จในแพลตฟอร์ม และป้องกันรถที่ไม่ได้ชาร์จครอบครองช่องจอด กำหนดงบโปรโมชั่นและวัดยอดใช้จริง โดยยังไม่นับรายได้ร้านค้าเพิ่มในผลตอบแทนนี้')
heading('ยอดใช้งานที่สัมพันธ์กับการคืนทุน')
table(['รถเข้าชาร์จเฉลี่ย', 'เงินสดดำเนินงานต่อปี บาท', 'ระยะคืนทุนโดยประมาณ'], [
    (f'{n} คันต่อวัน', money(model(n)['cash']), f"{model(n)['payback']:.2f} ปี")
    for n in (36, 44, 48)
], [5.2, 6.1, 5.9], 14.5)
para('เป้าคืนทุนภายใน 3 ปีต้องรักษายอดเฉลี่ยอย่างน้อยประมาณ 44 คันต่อวันภายใต้ต้นทุนเดิม เป้าหมาย 48 คันต่อวันจึงมีส่วนเผื่อ แต่ยังต้องพิสูจน์ยอดใช้บริการจริง', True, 15)
heading('ข้อมูลประกอบก่อนอนุมัติการลงทุน')
para('สถานี PluZ สุขุมเซอร์วิสแม่แตงอยู่ห่างประมาณ 1.1 กม. ในแนวตรง รายการสาธารณะระบุ DC 180 kW จึงควรแข่งขันด้วยการเข้าถึง ช่องจอด และบริการระหว่างพัก ไม่อ้างความเร็วเหนือคู่แข่งรายนี้ ส่วนตลาดแม่มาลัยอยู่ราว 1.3 กม. ไม่ใช่ร้านค้าภายในโครงการ', size=14)
para('สำรวจทางเข้าออกและจุดกลับรถ ตรวจโหลดและเงื่อนไขจ่ายไฟ พร้อมยืนยันขอบเขตเงินลงทุน 4.05 ล้านบาทก่อนสรุปราคา ต้นทุนไฟ 4.00 บาทต่อหน่วยเป็นสมมติฐานเฉลี่ยต่อหน่วยที่ขาย ไม่ใช่อัตราที่รับรองสำหรับพื้นที่นี้ ต้องตรวจค่าไฟรวมการสูญเสียและค่าใช้จ่ายจริงให้เพียงพอ', size=14)
para('ประมาณการไม่รวมภาษีเงินได้ ดอกเบี้ย เงินทุนหมุนเวียน ค่าใช้จ่ายนอก O&M และการเปลี่ยนอุปกรณ์ครั้งใหญ่ คงค่าใช้จ่ายเดิม 10 ปี ต้องทบทวนงบอินเทอร์เน็ต ประกันภัย และเงินสำรองซ่อมก่อนอนุมัติ ไม่ใช่การรับประกันผลตอบแทน', size=13.5)
link('ข้อมูลสถานที่ TAT STAR   Air Diamond Hotel', 'https://www.tatstar.org/company/STAR-B01601')
link('รายการสถานีชาร์จสาธารณะ EVNav TH   จังหวัดเชียงใหม่', 'https://evnavth.com/province/chiang-mai')
para('พงษ์วรา จันที   ผู้จัดการทั่วไป  TCE Project\nโทร. 0949953569   อีเมล pongwara.ja@tceproject.co.th', True, 14)

doc.core_properties.title = 'ข้อเสนอการลงทุนสถานีชาร์จรถยนต์ไฟฟ้า Air Diamond แม่แตง'
doc.core_properties.author = 'TCE Project'
doc.core_properties.subject = '180 kW จำนวน 2 ตู้ 4 ช่องจอด เป้าหมาย 48 คันต่อวัน'
path = OUT / 'Proposal_Air_Diamond_180kW_2Units_4Bays_TH_20261023.docx'
doc.save(path)
audit = dict(assumptions=dict(capex=str(CAPEX), cars_per_day=48, kwh_per_visit=str(KWH),
    selling_price=str(PRICE), electricity_effective_cost=str(POWER), cpo_rate=str(CPO),
    om_per_year=str(OM), rent=0, growth=0, demand_status='user_requested_commercial_target_not_validated'),
    annual={k: str(v) for k, v in m.items()},
    minimum_daily_for_3years=str((CAPEX / 3 + OM) / (KWH * 365 * (PRICE * (1 - CPO) - POWER))),
    ten_year_after_capex=str(m['cash'] * 10 - CAPEX))
(OUT / 'financial_audit.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')
assert m['cash'] == D('1509894.720')
assert CAPEX / m['cash'] < 3
assert m['cash'] * 10 - CAPEX == D('11048947.200')
print(path)
print(json.dumps(audit, ensure_ascii=False))
