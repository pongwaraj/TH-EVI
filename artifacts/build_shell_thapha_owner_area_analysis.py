"""Owner-facing Shell Tha Pha area analysis with explicit evidence boundaries."""

import json
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs/shell_thapha_owner_area_analysis"
OUT.mkdir(parents=True, exist_ok=True)
OUTPUT = OUT / "Area_Analysis_Shell_Tha_Pha_DC180kW_20261027.docx"
LOGO = ROOT / "artifacts/report_assets/tce_brand_logo.png"
FONT = "TH Sarabun New"
BLACK = RGBColor(0, 0, 0)
CAPEX = Decimal("2300500")


def money(value):
    return f"{Decimal(str(value)).quantize(Decimal('1'), rounding=ROUND_HALF_UP):,}"


def financial(cars):
    energy = Decimal(cars) * 35 * 365
    revenue = energy * Decimal("7.90")
    cpo = revenue * Decimal("0.08")
    electricity = energy * 4
    cash = revenue - cpo - electricity - 36000
    return {"cars_per_day": cars, "kwh_per_year": energy, "revenue_year": revenue,
            "revenue_month": revenue / 12, "electricity_year": electricity,
            "cpo_year": cpo, "om_year": Decimal(36000), "operating_cash_year": cash,
            "cumulative_cash_after_capex_year10": cash * 10 - CAPEX,
            "simple_payback_years": CAPEX / cash if cash > 0 else None}


def set_font(run, size=16, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = BLACK
    fonts = run._element.get_or_add_rPr().rFonts
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), FONT)


def paragraph(doc, text, *, size=16, bold=False, after=6, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.06
    p.paragraph_format.keep_with_next = keep
    set_font(p.add_run(text), size, bold)
    return p


def heading(doc, text, *, new_page=False):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = new_page
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run(text), 19, True)
    return p


def table(doc, headers, rows, widths, *, size=14):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    for index, cells in enumerate([headers] + list(rows)):
        row = t.rows[0] if index == 0 else t.add_row()
        row.height = None
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement("w:cantSplit")
        trpr.append(no_split)
        if index == 0:
            repeat = OxmlElement("w:tblHeader")
            trpr.append(repeat)
        for j, (cell, text) in enumerate(zip(row.cells, cells)):
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            if index == 0 or j > 0 and len(str(text)) < 28:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(str(text)), size, index == 0)
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for side, val in (("top", 55), ("bottom", 55), ("left", 90), ("right", 90)):
                element = OxmlElement("w:" + side)
                element.set(qn("w:w"), str(val))
                element.set(qn("w:type"), "dxa")
                margins.append(element)
            pr.append(margins)
            border = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                edge = OxmlElement("w:" + side)
                for key, val in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    edge.set(qn("w:" + key), val)
                border.append(edge)
            pr.append(border)
            if index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E8EDF2")
                pr.append(shade)
    paragraph(doc, "", size=5, after=1)
    return t


def source(doc, number, label, url):
    p = paragraph(doc, f"{number}. {label}  ", size=12, after=3)
    relationship = p.part.relate_to(url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT)
    fonts.set(qn("w:hAnsi"), FONT)
    props.append(fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "24")
    props.append(sz)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = "เปิดแหล่งข้อมูล"
    run.append(text)
    link.append(run)
    p._p.append(link)


def main():
    audit = json.loads((ROOT / "data/site_analysis/shell_thapha_thoen_20261027.json").read_text(encoding="utf-8"))
    click = audit["click_analysis"]
    assert click["aadt_used"] == 13337
    assert click["net_sessions_per_day"] == 2.0
    assert round(click["gross_area_demand_sessions"] - click["competitor_penalty_sessions"], 1) == 2.0
    cases = [financial(2)]
    projection = [{"operating_year": year, **cases[0],
                   "cumulative_operating_cash": cases[0]["operating_cash_year"] * year,
                   "cumulative_cash_after_capex": cases[0]["operating_cash_year"] * year - CAPEX}
                  for year in range(1, 11)]
    calculation = {"site": "Shell Tha Pha", "coordinates": [17.650561, 99.249117],
        "inputs": {"energy_kwh_per_visit_assumed": 35, "sale_thb_per_kwh_assumed": 7.9,
                   "electricity_thb_per_sold_kwh_assumed": 4, "cpo_share_assumed": .08,
                   "annual_om_assumed": 36000, "capex": CAPEX},
        "cases": cases, "projection": projection,
        "demand_basis": "legacy_screening_used_as_explicit_planning_assumption_not_observed_site_sales",
        "competitor_update_source": "Project coordinator supplied 2026-10-27",
        "competitor_updates": {"EleX PT Thoen": {"cabinet_kw": 124, "pin_available": False},
                               "EleX PT Thoen 3": {"cabinet_kw": 150},
                               "Spark Bangchak Thoen Napha": {"cabinet_kw": 240, "cabinets": 2}},
        "document_revision": "owner-facing single planning projection",
        "screening_model": {k: click.get(k) for k in ("gross_area_demand_sessions",
            "competitor_penalty_sessions", "net_sessions_per_day", "eligibility_status", "aadt_used")}}
    (OUT / "calculation_audit.json").write_text(json.dumps(calculation, ensure_ascii=False,
        indent=2, default=str), encoding="utf-8")

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(1.65), Cm(1.65)
    sec.left_margin, sec.right_margin = Cm(2), Cm(2)
    sec.header_distance, sec.footer_distance = Cm(.6), Cm(.65)
    for name, size in (("Normal", 16), ("Title", 25), ("Subtitle", 18), ("Heading 1", 19), ("Heading 2", 17)):
        style = doc.styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.paragraph_format.line_spacing = 1.06
        style.paragraph_format.space_after = Pt(6)
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + key), FONT)
    for style in doc.styles:
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(hp.add_run("TCE PROJECT   รายงานวิเคราะห์พื้นที่สถานีชาร์จรถยนต์ไฟฟ้า"), 11)
    fp = sec.footer.paragraphs[0]
    set_font(fp.add_run("Shell ท่าผา เถิน ลำปาง   •   27 ตุลาคม 2569"), 11)
    fp.add_run("\t")
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(16), WD_ALIGN_PARAGRAPH.RIGHT)
    set_font(fp.add_run("หน้า "), 11)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)
    doc.core_properties.title = "รายงานวิเคราะห์พื้นที่สถานีชาร์จรถยนต์ไฟฟ้า Shell ท่าผา"
    doc.core_properties.author = "TCE Project"
    doc.core_properties.subject = "Area Analysis for owner investment in one 180 kW DC charger"

    # Page 1: decision and evidence, not a sales forecast disguised as a heat score.
    if LOGO.exists():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(LOGO), width=Cm(2.6))
    p = doc.add_paragraph(style="Title")
    set_font(p.add_run("รายงานวิเคราะห์พื้นที่สถานีชาร์จรถยนต์ไฟฟ้า"), 25, True)
    paragraph(doc, "ปั๊มเชลล์ ท่าผา อำเภอเถิน จังหวัดลำปาง", size=21, bold=True)
    paragraph(doc, "DC Charger 180 kW จำนวน 1 ตู้ 2 ช่องจอด\nเงินลงทุนอ้างอิง 2,300,500 บาท สำหรับเจ้าของพื้นที่ลงทุนเอง", size=16)
    paragraph(doc, "ที่ตั้ง 197 หมู่ 5 ตำบลแม่ปะ อำเภอเถิน จังหวัดลำปาง 52160\nพิกัดวิเคราะห์ 17.650561, 99.249117", size=15)

    heading(doc, "ข้อสรุปสำหรับเจ้าของพื้นที่")
    paragraph(doc, "ทำเลมีโอกาสพัฒนาบริการชาร์จระหว่างการเดินทางและการแวะพัก โดยใช้พื้นที่และบริการเดิมของปั๊มเป็นฐาน ตู้ DC 180 kW จำนวน 1 ตู้ 2 ช่องจอดเป็นรูปแบบเริ่มต้นที่ให้บริการรถได้พร้อมกันสองคัน และไม่ต้องลงทุนหลายตู้ตั้งแต่วันแรก")
    paragraph(doc, "แนวทางของโครงการคือเพิ่มบริการชาร์จให้เป็นส่วนหนึ่งของการแวะพัก ลูกค้าสามารถชาร์จรถพร้อมใช้บริการภายในปั๊ม เจ้าของพื้นที่มีรายได้จากค่าชาร์จเพิ่มเติม และมีโอกาสเพิ่มยอดใช้บริการร้านค้าที่อยู่ในสถานี")
    table(doc, ["ข้อมูลพื้นที่", "ค่าที่ใช้", "ความหมายต่อโครงการ"], [
        ["รถผ่านแนวทางหลวง 1", "13,337 คัน/วัน", "AADT ปี 2568\nบริบทการเดินทางบนเส้นทางหลัก"],
        ["BEV รถยนต์นั่งส่วนบุคคลลำปาง", "1,725 คัน", "DLT ณ 30 ก.ย. 2569\nไม่รวมรถต่างจังหวัดที่ผ่านทาง"],
        ["สถานีชาร์จที่มีพิกัดใกล้เคียง", "ประมาณ 3.8–4.0 กม.", "มีการแข่งขันบนแนวถนน\nระยะในรายงานเป็นเส้นตรง"],
    ], [6.2, 3.7, 7.1], size=15)
    heading(doc, "จุดแข็งของโครงการ")
    paragraph(doc, "ส่วนเติมน้ำมันเปิดตลอด 24 ชั่วโมง และมีร้านกาแฟเป็นบริการประกอบ โครงการควรเน้นป้ายที่เห็นได้จากถนน ทางเข้าใช้ง่าย ช่องชาร์จว่างจริง และสถานะตู้ที่ตรวจผ่านแอปได้ หากจะเปิดชาร์จตลอดวัน ต้องจัดบริการสนับสนุนช่วงกลางคืนให้พร้อมด้วย")

    # Page 2: real basemap and nearby competition, with explicit source status.
    heading(doc, "บริบทพื้นที่และการแข่งขัน", new_page=True)
    map_path = OUT / "site_context_map.png"
    if not map_path.exists():
        raise FileNotFoundError("A verified real-basemap context map must be generated first")
    doc.add_picture(str(map_path), width=Cm(17))
    paragraph(doc, "S จุดวิเคราะห์   P1 โรงพยาบาลเถิน   P2 จุดขึ้นรถโดยสารเถิน\nC1–C3 กลุ่มสถานีชาร์จบริเวณเถิน   C4 Spark บางจาก เถิน–นภา", size=12, after=4)
    paragraph(doc, "แผนที่แสดงตำแหน่งเพื่อดูบริบท ไม่ใช่เส้นทางขับจริงหรือขอบเขตลูกค้าที่แน่นอน", size=12, after=7)
    table(doc, ["สถานีใกล้เคียง", "ระยะอ้างอิง", "กำลังชาร์จ"], [
        ["C1 EV Station เลขที่ 163 พหลโยธิน", "3.8 กม.", "DC CCS 120 kW"],
        ["C2 PluZ PTT ธ.1994 ขาเข้า", "3.9 กม.", "DC 120 kW จำนวน 2 หัว CCS2"],
        ["C3 EleX PT เถิน 3", "4.0 กม.", "DC 150 kW"],
        ["C4 Spark บางจาก เถิน–นภา", "9.0 กม.", "DC 240 kW จำนวน 2 ตู้"],
        ["EleX PT เถิน", "พื้นที่เถิน", "DC 124 kW"],
    ], [7.2, 2.2, 7.6], size=14)
    paragraph(doc, "ข้อมูลกำลังของ EleX PT เถิน EleX PT เถิน 3 และ Spark อัปเดตจากผู้ประสานงานโครงการ วันที่ 27 ตุลาคม 2569 สาขา PT เถินแสดงข้อมูลในตารางโดยไม่ระบุหมุด ส่วนระยะของ C1–C4 เป็นระยะเส้นตรงจากจุดวิเคราะห์", size=12)
    paragraph(doc, "สถานีของโครงการวางตำแหน่งบริการที่การแวะเข้าถึงง่าย ช่องจอดเฉพาะ และการใช้เวลาร่วมกับร้านกาแฟภายในปั๊ม ควบคู่กำลังชาร์จ 180 kW เพื่อเป็นทางเลือกให้ผู้เดินทางบนเส้นทางนี้", size=15)

    # Page 3: equipment and demand interpretation without treating screening as sales.
    heading(doc, "รูปแบบสถานีที่ใช้ในการวิเคราะห์", new_page=True)
    table(doc, ["รายการ", "รูปแบบเสนอพิจารณา"], [
        ["ตู้ชาร์จ DC", "กำลังรวม 180 kW จำนวน 1 ตู้"],
        ["หัวชาร์จและช่องจอด", "CCS2 จำนวน 2 หัว พร้อมช่องจอดเฉพาะ 2 ช่อง"],
        ["การแบ่งกำลัง", "กำลังรวมทั้งตู้ไม่เกิน 180 kW\nเมื่อชาร์จสองคันให้แบ่งกำลังตามระบบที่เลือก"],
        ["เวลาเปิดบริการ", "วางแผนบริการ 24 ชั่วโมง พร้อมระบบช่วยเหลือลูกค้า"],
    ], [5.1, 11.9], size=15)
    paragraph(doc, "ขนาดที่เสนอเหมาะกับการเริ่มให้บริการเป็นจุดพักริมทาง และรองรับการใช้งานพร้อมกันสองคัน ควรเลือกการบริหารกำลังแบบปรับตามความต้องการของรถ พร้อมเตรียมแนวสายไฟและพื้นที่สำหรับขยายในอนาคต โดยยังไม่จำเป็นต้องติดตั้งตู้ที่สองก่อนทราบยอดใช้บริการจริง")
    paragraph(doc, "ระบบแบ่งกำลังให้เหมาะกับรถที่เข้าใช้บริการและระดับแบตเตอรี่ สำหรับการวางแผนช่องจอด หากรับไฟ 35 kWh ที่กำลังเฉลี่ย 60 kW จะใช้เวลาชาร์จประมาณ 35 นาที และเผื่อเข้าออกช่องจอดอีก 5 นาที รวมประมาณ 40 นาที โดยเวลาชาร์จจริงเปลี่ยนตามรถและการใช้งานพร้อมกัน", size=15)

    heading(doc, "ประมาณการใช้บริการ")
    paragraph(doc, "รายงานใช้ประมาณการเริ่มต้นเฉลี่ย 2 คันต่อวัน หลังพิจารณาสภาพพื้นที่และการแข่งขัน โดยกำหนดพลังงานเฉลี่ย 35 kWh ต่อการเข้าชาร์จ คิดเป็นการจำหน่ายไฟประมาณ 70 kWh ต่อวัน ตัวเลขนี้เป็นสมมติฐานวางแผนจากผลวิเคราะห์พื้นที่ ไม่ใช่ยอดใช้บริการที่สำรวจจากสถานีนี้")
    paragraph(doc, "กลุ่มลูกค้าเป้าหมายคือผู้ใช้รถ EV ที่เดินทางบนทางหลวง 1 และต้องการแวะพักระหว่างทาง รวมถึงผู้ใช้รถในพื้นที่เถิน ข้อมูลรถจดทะเบียนใช้ประกอบภาพตลาดท้องถิ่น ส่วนรถผ่านทางหลวงใช้ประกอบโอกาสเข้าถึงผู้เดินทาง โดยแยกจากจำนวนผู้เข้าชาร์จในประมาณการ")
    heading(doc, "แนวทางดำเนินโครงการ")
    paragraph(doc, "จัดช่องจอดเฉพาะสำหรับชาร์จ 2 ช่อง ติดป้ายก่อนถึงทางเข้า และออกแบบจุดจอดให้เห็นได้ชัดจากภายในปั๊ม เตรียมระบบไฟฟ้า ระบบชำระเงิน และช่องทางช่วยเหลือลูกค้าให้สอดคล้องกับเวลาบริการ พร้อมแสดงสถานีในแอปและ Google Maps")
    paragraph(doc, "หลังเปิดบริการ ให้ติดตามจำนวนรถ ปริมาณไฟที่จำหน่าย และช่วงเวลาการใช้งาน เพื่อปรับการประชาสัมพันธ์ร่วมกับร้านค้าและวางแผนขยายสถานีตามยอดใช้บริการ", size=15)

    # Page 4: conditional returns at the supplied investment, not a sales forecast.
    heading(doc, "ประมาณการรายได้และเงินสด 10 ปี", new_page=True)
    paragraph(doc, "จากการใช้บริการเฉลี่ย 2 คันต่อวัน รายได้ประมาณ 16,820 บาทต่อเดือน หรือ 201,845 บาทต่อปี หลังหักค่าไฟ ค่าระบบ CPO และค่าใช้จ่ายประจำ มีเงินสดจากการดำเนินงานประมาณ 47,497 บาทต่อปี", size=15)
    paragraph(doc, "ใช้เงินลงทุนเริ่มต้น 2,300,500 บาท คิดปีละ 365 วัน โดยคงจำนวนรถ ราคาขาย และต้นทุนเดิมตลอด 10 ปี ตารางแสดงปีสำคัญนับจากวันเปิดบริการ และเงินสดสะสมคอลัมน์สุดท้ายหักเงินลงทุนเริ่มต้นแล้ว", size=14)
    table(doc, ["ปีดำเนินงาน", "รถชาร์จ\nคัน/วัน", "รายได้\nบาท/ปี", "เงินสดดำเนินงาน\nบาท/ปี", "เงินสดสะสม\nหลังลงทุน บาท"], [
        [str(r["operating_year"]), str(r["cars_per_day"]), money(r["revenue_year"]),
         money(r["operating_cash_year"]), money(r["cumulative_cash_after_capex"])]
        for r in projection if r["operating_year"] in (1, 3, 5, 10)
    ], [2.1, 2.1, 3.6, 4.3, 4.9], size=13)
    paragraph(doc, "รายได้รวม 10 ปีประมาณ 2,018,450 บาท และเงินสดดำเนินงานสะสม 474,974 บาท ภายใต้ประมาณการนี้ เงินลงทุนยังไม่ถูกชดเชยครบภายใน 10 ปี", size=14)
    paragraph(doc, "สมมติฐานเพื่อเปรียบเทียบ ใช้ 35 kWh ต่อคัน ราคาขาย 7.90 บาทต่อ kWh ต้นทุนไฟฟ้ารวมผลสูญเสียเฉลี่ย 4.00 บาทต่อหน่วยที่ขาย ค่าระบบ CPO 8% ของรายได้ และค่าดำเนินงานกับบำรุงรักษา 36,000 บาทต่อปี คิด 365 วันต่อปี ใช้ค่าที่ยังไม่ปัดเศษในการคำนวณและแสดงจำนวนเงินเป็นบาท", size=13)
    paragraph(doc, "เงินสดดำเนินงานเป็นรายรับหลังหักค่าไฟ CPO และค่าใช้จ่ายประจำ แต่ก่อนเงินลงทุน ภาษี ดอกเบี้ย และค่าเปลี่ยนอุปกรณ์ครั้งใหญ่ ไม่ใช่กำไรสุทธิ ไม่ใส่ค่าเช่าหรือ GP เจ้าของที่ดิน โดยสมมติว่าไม่มีภาระดังกล่าว คันต่อวันหมายถึงการเข้าใช้บริการชาร์จ ไม่ใช่จำนวนทะเบียนรถไม่ซ้ำ", size=13)
    heading(doc, "แหล่งข้อมูลสำคัญ")
    for number, label, url in [
        (1, "Shell locator ยืนยันสาขาและเวลาบริการเติมน้ำมัน", "https://find.shell.com/th/fuel/12849483-do-subcharean-oil-thoen-lpg"),
        (2, "กรมทางหลวง AADT ปี 2568 ตอน 1102 กม 617+097", "https://datagov.mot.go.th/dataset/traf62"),
        (3, "DLT ยอดสะสม 30 ก.ย. 2569 รย.1 ลำปาง", "https://web.dlt.go.th/statistics/load_file_select_new_car.php?t=3&data_file=1302"),
        (4, "DOPA ปี 2568 เถินรวมเทศบาล 55,218 คน", "https://stat.bora.dopa.go.th/new_stat/webPage/statByYear.php"),
        (5, "EGAT รายชื่อสถานี EleX สาขาเถิน", "https://egatev.egat.co.th/our-services/elex-by-egat/"),
        (6, "EVNav รายการ PluZ ลำปาง อัปเดต 25 ก.ย. 2569", "https://evnavth.com/province/lampang"),
    ]:
        source(doc, number, label, url)
    paragraph(doc, "ข้อมูลกำลังตู้คู่แข่งปรับตามข้อมูลผู้ประสานงานโครงการ ณ 27 ตุลาคม 2569 ประมาณการรายได้ครอบคลุมค่าชาร์จเท่านั้น แยกจากรายได้ส่วนเพิ่มของร้านค้าและบริการภายในปั๊ม", size=12)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
